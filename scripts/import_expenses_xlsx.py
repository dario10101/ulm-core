"""Importa gastos exportados de la app de finanzas anterior (archivos .xlsx).

Lee solo la hoja "Expenses". Fila 1: rango del export (se ignora); fila 2:
encabezados; desde la fila 3, un gasto por fila:

    A Date and time   -> recorded_on (solo la fecha; la hora siempre es 00:00)
    B Category        -> category_id, por nombre exacto contra fn_categories
    C Account         -> payment_method_id (ver ACCOUNT_TO_PAYMENT_METHOD)
    D Amount in default currency -> amount (COP; F trae el mismo valor)
    K Comment         -> name, en mayusculas como lo guarda la API. Si viene
                         vacio se usa el nombre de la categoria.

El resto de columnas no se usa (moneda siempre COP, tags siempre vacios).

Valida TODO antes de escribir: si una categoria o cuenta no calza, aborta
sin insertar nada. Para no duplicar, se niega a correr si el usuario ya tiene
gastos en el rango de fechas de los archivos (--force lo salta).

Uso (desde ulm-core/, con el venv activado y openpyxl de requirements-dev.txt):
    python -m scripts.import_expenses_xlsx ../ulm-data/expenses/*.xlsx \
        --email ruben.d21pc@gmail.com --dry-run
    python -m scripts.import_expenses_xlsx ../ulm-data/expenses/*.xlsx \
        --email ruben.d21pc@gmail.com
"""

import argparse
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select

import app.main  # noqa: F401 -- registra todos los modelos
from app.db.models.finance import Category, Expense, PaymentMethod
from app.db.models.user import User
from app.db.session import SessionLocal

SHEET = "Expenses"
EXPECTED_HEADER = ("Date and time", "Category", "Account", "Amount in default currency")

# Cuenta de la app anterior -> nombre del metodo de pago en fn_payment_methods.
ACCOUNT_TO_PAYMENT_METHOD = {
    "Cash": "Cash",
    "Credit cards": "Credit card",
    "Nequi": "Bank transfer",
    "Bancolombia": "Bank transfer",
}


@dataclass(frozen=True)
class Row:
    source: str
    recorded_on: date
    category: str
    account: str
    amount: Decimal
    comment: str


def read_rows(path: Path) -> list[Row]:
    # Sin read_only: estos exports no traen la dimension de la hoja y en ese
    # modo openpyxl no devuelve filas.
    sheet = load_workbook(path, data_only=True)[SHEET]
    values = list(sheet.iter_rows(values_only=True))
    header = tuple(values[1][: len(EXPECTED_HEADER)])
    if header != EXPECTED_HEADER:
        sys.exit(f"{path.name}: encabezados inesperados en la fila 2: {header}")

    rows = []
    for line, cells in enumerate(values[2:], start=3):
        if all(cell in (None, "") for cell in cells):
            continue
        when, category, account, amount = cells[0], cells[1], cells[2], cells[3]
        comment = cells[10]
        if when is None or not category or not account or amount is None:
            sys.exit(f"{path.name}:{line}: fila incompleta: {cells}")
        rows.append(
            Row(
                source=f"{path.name}:{line}",
                recorded_on=when.date(),
                category=category.strip(),
                account=account.strip(),
                amount=Decimal(str(amount)),
                comment=(comment or "").strip(),
            )
        )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--email", required=True, help="dueño de los gastos")
    parser.add_argument("--dry-run", action="store_true", help="valida y resume, no inserta")
    parser.add_argument("--force", action="store_true", help="inserta aunque ya haya gastos")
    args = parser.parse_args()

    rows = [row for path in args.files for row in read_rows(path)]
    if not rows:
        sys.exit("No hay filas para importar.")

    with SessionLocal() as db:
        user = db.scalar(select(User).where(func.lower(User.email) == args.email.lower()))
        if user is None:
            sys.exit(f"No existe el usuario {args.email}.")

        categories = {c.name: c.id for c in db.scalars(select(Category))}
        payment_methods = {p.name: p.id for p in db.scalars(select(PaymentMethod))}

        # Validacion completa antes de escribir nada.
        errors = []
        for row in rows:
            if row.category not in categories:
                errors.append(f"{row.source}: categoria desconocida {row.category!r}")
            method = ACCOUNT_TO_PAYMENT_METHOD.get(row.account)
            if method is None or method not in payment_methods:
                errors.append(f"{row.source}: cuenta sin mapear {row.account!r}")
            if row.amount <= 0:
                errors.append(f"{row.source}: monto invalido {row.amount}")
        if errors:
            sys.exit("Nada insertado:\n" + "\n".join(errors[:50]))

        first = min(row.recorded_on for row in rows)
        last = max(row.recorded_on for row in rows)
        existing = db.scalar(
            select(func.count())
            .select_from(Expense)
            .where(
                Expense.user_id == user.id,
                Expense.recorded_on.between(first, last),
            )
        )
        if existing and not args.force:
            sys.exit(
                f"{user.email} ya tiene {existing} gastos entre {first} y {last}; "
                "para no duplicar no se importa nada (usa --force si es a proposito)."
            )

        total = sum(row.amount for row in rows)
        print(f"{len(rows)} gastos de {first} a {last}, total $ {total:,.0f} -> {user.email}")
        print("  por cuenta:", dict(Counter(row.account for row in rows)))
        print("  sin comentario (name = categoria):", sum(1 for row in rows if not row.comment))
        if args.dry_run:
            print("--dry-run: no se inserto nada.")
            return

        db.add_all(
            Expense(
                user_id=user.id,
                name=(row.comment or row.category).upper()[:200],
                amount=row.amount,
                recorded_on=row.recorded_on,
                note=None,
                category_id=categories[row.category],
                payment_method_id=payment_methods[ACCOUNT_TO_PAYMENT_METHOD[row.account]],
                tags=[],
            )
            for row in rows
        )
        db.commit()  # una sola transaccion: o entran todos o ninguno
        print("Importados.")


if __name__ == "__main__":
    main()
